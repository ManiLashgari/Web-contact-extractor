"""
Utilities for saving extracted website contacts to Excel files.
"""

import os
from openpyxl import Workbook, load_workbook


def save_contacts(
    excel_path: str,
    website_url: str,
    contacts: dict[str, list[str]],
) -> None:
    """
    Save extracted website contacts to an Excel workbook without duplicates.

    Creates a new workbook with column headers if the file does not exist.
    Otherwise, adds only contact records that are not already present.

    Parameters
    ----------
    excel_path : str
        Path to the Excel workbook.
        The parent directory is created if it does not exist.

    website_url : str
        URL of the website associated with the extracted contacts.

    contacts : dict[str, list[str]]
        Dictionary mapping contact types to lists of extracted values.
        Each contact is stored as a separate record.

    Returns
    -------
    None
        Saves the workbook to the specified path.

    Notes
    -----
    The workbook contains a worksheet named ``Contacts`` with three columns:

    - URL: Website URL associated with the contact.
    - Type: Type of contact, such as Email or Phone number.
    - Value: Extracted contact information.

    Duplicate records are identified by comparing the website URL,
    contact type, and contact value.
    """
    os.makedirs(os.path.dirname(excel_path) or ".", exist_ok=True)

    if os.path.exists(excel_path):
        workbook = load_workbook(excel_path)
        if "Contacts" in workbook.sheetnames:
            worksheet = workbook["Contacts"]
        else:
            worksheet = workbook.create_sheet("Contacts")
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Contacts"

    # Add headers if the first row is empty.
    if all(
        worksheet.cell(row=1, column=col).value is None
        for col in range(1, 4)
    ):
        worksheet.cell(row=1, column=1, value="URL")
        worksheet.cell(row=1, column=2, value="Type")
        worksheet.cell(row=1, column=3, value="Value")

    # Collect existing contact records, starting after the header.
    existing_contacts = {
        (row[0], row[1], row[2])
        for row in worksheet.iter_rows(
            min_row=2,
            max_col=3,
            values_only=True,
        )
        if all(value is not None for value in row)
    }

    # Insert only new contacts.
    for contact_type, contact_list in contacts.items():
        for contact in contact_list:
            contact_record = (website_url, contact_type, contact)

            if contact_record not in existing_contacts:
                worksheet.append(list(contact_record))
                existing_contacts.add(contact_record)

    workbook.save(excel_path)
