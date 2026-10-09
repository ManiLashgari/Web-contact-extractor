"""
Unit tests for saving website contacts to Excel files.
"""

from openpyxl import load_workbook

from src.excel import save_contacts


def test_save_contacts_creates_excel_file(tmp_path):
    """
    Verify that an Excel file is created with the expected headers.
    """
    excel_path = tmp_path / "contacts.xlsx"

    save_contacts(
        str(excel_path),
        "https://example.com",
        ["hello@example.com"],
    )

    assert excel_path.exists()

    workbook = load_workbook(excel_path)
    worksheet = workbook["Contacts"]

    assert worksheet["A1"].value == "Website"
    assert worksheet["B1"].value == "Type"
    assert worksheet["C1"].value == "Value"


def test_save_contacts_writes_email_data(tmp_path):
    """
    Verify that extracted emails are written to the worksheet.
    """
    excel_path = tmp_path / "contacts.xlsx"

    save_contacts(
        str(excel_path),
        "https://example.com",
        ["hello@example.com", "support@example.com"],
    )

    workbook = load_workbook(excel_path)
    worksheet = workbook["Contacts"]

    assert worksheet["A2"].value == "https://example.com"
    assert worksheet["B2"].value == "Email"
    assert worksheet["C2"].value == "hello@example.com"

    assert worksheet["A3"].value == "https://example.com"
    assert worksheet["B3"].value == "Email"
    assert worksheet["C3"].value == "support@example.com"


def test_save_contacts_with_empty_data(tmp_path):
    """
    Verify that an Excel file with headers is created when no emails exist.
    """
    excel_path = tmp_path / "contacts.xlsx"

    save_contacts(str(excel_path), "https://example.com", [])

    workbook = load_workbook(excel_path)
    worksheet = workbook["Contacts"]

    assert worksheet.max_row == 1
    assert [cell.value for cell in worksheet[1]] == [
        "Website",
        "Type",
        "Value",
    ]
